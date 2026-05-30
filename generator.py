import os
import io
import json
import pdfplumber
import fitz 
from PIL import Image
import pytesseract
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from openai import OpenAI

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
OPENROUTER_API_KEY = "your_openrouter_api_key_here"
client = OpenAI(
    base_url="https://openrouter.ai/api/v1",
    api_key=OPENROUTER_API_KEY,
)
MODEL_NAME = "openrouter/free"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_DIR = os.path.join(SCRIPT_DIR, "inputs")
OUTPUT_DIR = os.path.join(SCRIPT_DIR, "outputs")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "A2_Vocabulary_Book.docx")
def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            if not text.strip():
                print(f"PDF appears to be scanned. Converting pages to images for OCR...")
                doc = fitz.open(file_path)
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    pix = page.get_pixmap(dpi=300)  # Render page to high-res image
                    img_data = pix.tobytes("png")
                    img = Image.open(io.BytesIO(img_data))
                    text += pytesseract.image_to_string(img) + "\n"
                doc.close()
                    
        elif ext in [".png", ".jpg", ".jpeg"]:
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            
    except Exception as e:
        print(f"Error reading file {os.path.basename(file_path)}: {e}")
    return text
def get_vocabulary_data(raw_text, unit_name):
    prompt = f"""
    Analyze the text provided below from '{unit_name}'. 
    Identify and extract vocabulary words that fit the A2 to C2 (A2+) CEFR level.
    Exclude very basic words (A1).
    
    For each word, provide:
    1. The English word
    2. Part of speech (e.g., n., v., adj.)
    3. A simple English definition
    4. Russian translation
    5. Uzbek translation
    6. A simple English example sentence using the word
    
    Return the result strictly as a JSON object containing a list named "words", formatted as follows:
    {{
      "words": [
        {{
          "word": "example",
          "part_of_speech": "n.",
          "definition": "a representative form",
          "translation_ru": "пример",
          "translation_uz": "misol",
          "example": "This is a clear example."
        }}
      ]
    }}
    Do not add any conversational text, formatting notes, or markdown code blocks (like ```json). Return raw JSON only.

    Text to analyze:
    {raw_text}
    """
    
    try:
        response = client.chat.completions.create(
            model=MODEL_NAME,
            messages=[{"role": "user", "content": prompt}],
            extra_headers={
                "HTTP-Referer": "https://localhost:3000",
                "X-Title": "Vocab Builder Tool",
            }
        )
        
        raw_content = response.choices[0].message.content.strip()
        
        if raw_content.startswith("```json"):
            raw_content = raw_content[7:]
        elif raw_content.startswith("```"):
            raw_content = raw_content[3:]
        if raw_content.endswith("```"):
            raw_content = raw_content[:-3]
        raw_content = raw_content.strip()
        
        data = json.loads(raw_content)
        return data.get("words", [])
    except Exception as e:
        print(f"Failed to process API request for {unit_name}: {e}")
        return []

def setup_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
      
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    return doc

def add_title_page(doc):
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("\n\n\n\n\n\nA2+ VOCABULARY BOOK")
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 78, 121)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Organized by Units with Definitions and Translations")
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(120, 120, 120)
    
    doc.add_page_break()


def add_unit_table(doc, unit_title, words):
    heading = doc.add_heading(level=1)
    heading_run = heading.add_run(unit_title)
    heading_run.font.color.rgb = RGBColor(31, 78, 121)
    heading_run.font.size = Pt(18)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(12)
    
    if not words:
        doc.add_paragraph("No A2+ vocabulary extracted for this unit.")
        return
    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    table.autofit = False
    
    col_widths = [Inches(1.0), Inches(0.5), Inches(1.5), Inches(1.0), Inches(1.0), Inches(1.5)]
    headers = ["Word", "P.O.S.", "Definition", "Russian", "Uzbek", "Example"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1F4E79")
        
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    for word_data in words:
        row_cells = table.add_row().cells
        
        row_cells[0].text = word_data.get("word", "")
        row_cells[1].text = word_data.get("part_of_speech", "")
        row_cells[2].text = word_data.get("definition", "")
        row_cells[3].text = word_data.get("translation_ru", "")
        row_cells[4].text = word_data.get("translation_uz", "")
        row_cells[5].text = word_data.get("example", "")
        
        for idx, cell in enumerate(row_cells):
            cell.width = col_widths[idx]
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.space_after = Pt(6)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    if idx == 0:
                        run.font.bold = True 


def main():
    if not os.path.exists(OUTPUT_DIR):
        os.makedirs(OUTPUT_DIR)
        
    doc = setup_document()
    add_title_page(doc)
    
    if not os.path.exists(INPUT_DIR):
        os.makedirs(INPUT_DIR)
        print(f"Created '{INPUT_DIR}' directory. Please add your files and run the script again.")
        return
        
    files = sorted([f for f in os.listdir(INPUT_DIR) if os.path.isfile(os.path.join(INPUT_DIR, f))])
    
    if not files:
        print("The 'inputs' folder is empty. Place your files there before running the script.")
        return
        
    for index, filename in enumerate(files):
        file_path = os.path.join(INPUT_DIR, filename)
        
        clean_name = os.path.splitext(filename)[0].replace("_", " ").title()
        unit_title = f"Unit {index + 1}: {clean_name}"
        
        print(f"Processing {unit_title} from {filename}...")
        
        raw_text = extract_text(file_path)
        if not raw_text.strip():
            print(f"Skipping empty or unreadable file: {filename}")
            continue
            
        words = get_vocabulary_data(raw_text, unit_title)
        add_unit_table(doc, unit_title, words)
        if index < len(files) - 1:
            doc.add_page_break()
            
    doc.save(OUTPUT_FILE)
    print(f"\nDocument generation complete. Saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
